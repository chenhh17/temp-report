from __future__ import annotations

import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))
from fill_final_report import (
    read_sections,
    split_milestones,
    populate_cell,
    set_cell_text,
)

TEMPLATE = ROOT / "audit_example" / "[company name]_Final Report Template.docx"
SOURCE = ROOT / "TadRreamk_Final Report.md"
OUTPUT = ROOT / "TadReamk Limited Final Report (draft).docx"

REPORTING_PERIOD = "[1 April 2025 to 31 March 2026]"
HEADER_SUBMITTED = "Submitted by TadReamk Limited on [Date]"

SECTION_MAP = {
    0: "Executive summary of the progress",
    1: "Deliverables / technological and financial achievements",
    3: "Progress of premise rental, equipment rental / purchase and manpower expenditure",
    4: "Any potential collaborations / M&A / investors",
    5: "Difficulties encountered",
    6: "Outstanding issues (if any)",
    7: "Innovation and technology content and commercialisation",
    8: "Commercial viability of the business",
    9: "Capability of your technology start-up and your team to undertake the R&D work and manage the company",
    10: "Social and/or community impacts of the technology start-up's R&D work",
}


def build() -> None:
    sections = read_sections(SOURCE)

    import shutil
    shutil.copy2(TEMPLATE, OUTPUT)
    doc = Document(OUTPUT)

    doc.paragraphs[0].text = f"TSSSU Final Report (Reporting period: {REPORTING_PERIOD} )"
    doc.paragraphs[1].text = HEADER_SUBMITTED

    missing = [k for k in SECTION_MAP.values() if k not in sections]
    if missing:
        raise RuntimeError(f"Missing sections in markdown: {missing}")

    for table_idx, key in SECTION_MAP.items():
        populate_cell(doc.tables[table_idx].cell(0, 0), sections[key])

    milestone_rows = split_milestones(sections.get("Milestones achieved", ""))
    milestone_table = doc.tables[2]
    while len(milestone_table.rows) < len(milestone_rows) + 2:
        milestone_table.add_row()
    while len(milestone_table.rows) > len(milestone_rows) + 2:
        tr = milestone_table.rows[-1]._tr
        milestone_table._tbl.remove(tr)

    for idx, (date_from, date_to, milestone, achieved) in enumerate(milestone_rows):
        row_idx = idx + 2
        set_cell_text(milestone_table, row_idx, 0, date_from)
        set_cell_text(milestone_table, row_idx, 1, date_to)
        set_cell_text(milestone_table, row_idx, 2, milestone)
        set_cell_text(milestone_table, row_idx, 3, achieved)

    doc.save(OUTPUT)
    print(f"Built: {OUTPUT}")
    print(f"Milestone rows: {len(milestone_rows)}")


if __name__ == "__main__":
    build()
