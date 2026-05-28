from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent
TEMPLATE = ROOT / "[company name]_Final Report Template.docx"
OUTPUT = ROOT / "TadReamk Final Report (draft).docx"
SOURCE = ROOT / "TadRreamk_Final Report.md"

COMPANY = "TadReamk Limited"
SUBMIT_DATE = "29 May 2026"
REPORTING_PERIOD = "1 April 2025 to 31 March 2026"


def read_sections(md_path: Path) -> dict[str, str]:
    text = md_path.read_text(encoding="utf-8")
    pattern = re.compile(r"^## (\d+)\.\s+(.+)$", re.MULTILINE)
    matches = list(pattern.finditer(text))
    sections: dict[str, str] = {}
    for idx, match in enumerate(matches):
        key = match.group(2).strip()
        start = match.end()
        end = matches[idx + 1].start() if idx + 1 < len(matches) else len(text)
        body = text[start:end].strip()
        body = re.sub(r"^---+\s*", "", body)
        body = re.sub(r"\s*---+\s*$", "", body)
        sections[key] = body.strip()
    return sections


def split_milestones(section_text: str) -> list[tuple[str, str, str, str]]:
    rows: list[tuple[str, str, str, str]] = []
    for line in section_text.splitlines():
        line = line.strip()
        if not line.startswith("|"):
            continue
        if "Period" in line or "---" in line:
            continue
        parts = [p.strip() for p in line.strip("|").split("|")]
        if len(parts) != 4:
            continue
        rows.append(tuple(parts))
    return rows


def add_runs_with_bold(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            italic_parts = re.split(r"(\*[^*]+\*)", part)
            for sub in italic_parts:
                if not sub:
                    continue
                if sub.startswith("*") and sub.endswith("*") and not sub.startswith("**"):
                    run = paragraph.add_run(sub[1:-1])
                    run.italic = True
                else:
                    paragraph.add_run(sub)


def normalize_markdown(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = text.replace("_", "")
    return text


def populate_cell(cell, content: str) -> None:
    cell.text = ""
    content = normalize_markdown(content)
    lines = content.splitlines()
    first = True
    for raw_line in lines:
        line = raw_line.rstrip()
        if not line.strip():
            if not first:
                cell.add_paragraph("")
            continue

        bullet_match = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        number_match = re.match(r"^(\d+)\.\s+(.*)$", line)
        heading_match = re.match(r"^###\s+(.*)$", line)

        if first:
            paragraph = cell.paragraphs[0]
            first = False
        else:
            paragraph = cell.add_paragraph()

        if heading_match:
            run = paragraph.add_run(heading_match.group(1).strip())
            run.bold = True
            continue

        if bullet_match:
            paragraph.style = "List Paragraph"
            add_runs_with_bold(paragraph, bullet_match.group(2).strip())
            continue

        if number_match:
            paragraph.style = "List Paragraph"
            add_runs_with_bold(
                paragraph, f"{number_match.group(1)}. {number_match.group(2).strip()}"
            )
            continue

        if line.startswith(">"):
            add_runs_with_bold(paragraph, line.lstrip("> ").strip())
            continue

        add_runs_with_bold(paragraph, line.strip())


def set_cell_text(table, row: int, col: int, text: str) -> None:
    cell = table.cell(row, col)
    cell.text = ""
    paragraph = cell.paragraphs[0]
    add_runs_with_bold(paragraph, normalize_markdown(text))


def fill_document() -> None:
    sections = read_sections(SOURCE)

    section_map = {
        0: "Executive summary of the progress",
        1: "Deliverables / technological and financial achievements",
        3: "Progress of premise rental, equipment rental / purchase and manpower expenditure",
        4: "Any potential collaborations / M&A / investors",
        5: "Difficulties encountered",
        6: "Outstanding issues (if any)",
        7: "Innovation and technology content and commercialisation",
        8: "Commercial viability of the business",
        9: "Capability of your technology start-up and your team to undertake the R&D work and manage the company",
        10: "Social and/or community impacts of the technology start-up's R&D work",
    }

    shutil.copy2(TEMPLATE, OUTPUT)
    doc = Document(OUTPUT)

    doc.paragraphs[0].text = f"TSSSU Final Report (Reporting period: {REPORTING_PERIOD})"
    doc.paragraphs[1].text = f"Submitted by {COMPANY} on {SUBMIT_DATE}"

    for table_idx, section_key in section_map.items():
        content = sections.get(section_key, "")
        populate_cell(doc.tables[table_idx].cell(0, 0), content)

    milestone_rows = split_milestones(sections.get("Milestones achieved", ""))
    milestone_table = doc.tables[2]
    while len(milestone_table.rows) < len(milestone_rows) + 2:
        milestone_table.add_row()

    for idx, (date_from, date_to, milestone, achieved) in enumerate(milestone_rows):
        row_idx = idx + 2
        set_cell_text(milestone_table, row_idx, 0, date_from)
        set_cell_text(milestone_table, row_idx, 1, date_to)
        set_cell_text(milestone_table, row_idx, 2, milestone)
        set_cell_text(milestone_table, row_idx, 3, achieved)

    doc.save(OUTPUT)


if __name__ == "__main__":
    fill_document()
    print(f"Created: {OUTPUT}")
